#!/usr/bin/env python3
"""Build the course PowerPoint template: slides/template/msacl_ds301.pptx.

This file is the single home of the deck look (PPTX-first toolchain, see
course_plan/DECISIONS.md 2026-09-08). It carries a real Office theme (course
palette + fonts), a styled slide master, and six purpose-built layouts:

    Title Slide          eyebrow · big title · teal subtitle · amber bar
    Title and Body       eyebrow · 34pt title · bulleted body (teal bullets)
    Two Column           eyebrow · title · two content boxes
    Screenshot + Caption eyebrow · title · picture placeholder · muted caption
    Title Only           eyebrow · title, rest free
    Blank                paper background only

Anyone editing a deck in PowerPoint picks New Slide -> layout and everything
is already on-style; theme colors put the palette in every color picker.
tools/adopt_template.py splices this master into existing decks;
tools/spec2pptx.py scaffolds new structured slides from it.

Usage:
    python tools/make_template.py          # (re)build the template
"""

import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "slides" / "template" / "msacl_ds301.pptx"

# ---- course palette (was: constants inside spec2pptx.py) --------------------
PAPER = "FBFBF8"       # background            -> theme lt1 / bg1
INK = "1B2025"         # primary text          -> theme dk1 / tx1
INK_SOFT = "3D444C"    # secondary text        -> theme dk2 / tx2
SIGNAL_SOFT = "E3F0EF" # teal wash             -> theme lt2 / bg2
SIGNAL = "0E7C7B"      # teal accent           -> accent1
ROI = "E09E2F"         # amber accent          -> accent2
NEG = "C4534F"         # negative red          -> accent3
MUTED = "8A9099"       # muted grey            -> accent4
ROI_INK = "A9721A"     # dark amber            -> accent5
HAIRLINE = "D9D9D2"    # hairline grey         -> accent6

SANS = "Avenir Next"
MONO = "Menlo"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
M = Inches(0.6)          # page margin
USABLE = Inches(12.13)   # 13.333 - 2*0.6

KEEP = {  # default-template layout name -> our layout name
    "Title Slide": "Title Slide",
    "Title and Content": "Title and Body",
    "Two Content": "Two Column",
    "Picture with Caption": "Screenshot + Caption",
    "Title Only": "Title Only",
    "Blank": "Blank",
}

A = nsdecls("a")
EYEBROW_IDX = "13"


# ---- theme -------------------------------------------------------------------
def restyle_theme(theme_part):
    """Rewrite the theme's color scheme and font scheme to the course look."""
    root = etree.fromstring(theme_part.blob)
    root.set("name", "MSACL DS301")
    scheme = root.find(f".//{qn('a:clrScheme')}")
    scheme.set("name", "MSACL DS301")
    colors = {
        "dk1": INK, "lt1": PAPER, "dk2": INK_SOFT, "lt2": SIGNAL_SOFT,
        "accent1": SIGNAL, "accent2": ROI, "accent3": NEG, "accent4": MUTED,
        "accent5": ROI_INK, "accent6": HAIRLINE, "hlink": SIGNAL,
        "folHlink": ROI_INK,
    }
    for tag, val in colors.items():
        el = scheme.find(qn(f"a:{tag}"))
        for child in list(el):
            el.remove(child)
        srgb = etree.SubElement(el, qn("a:srgbClr"))
        srgb.set("val", val)
    fonts = root.find(f".//{qn('a:fontScheme')}")
    fonts.set("name", "MSACL DS301")
    for group in (qn("a:majorFont"), qn("a:minorFont")):
        latin = fonts.find(group).find(qn("a:latin"))
        latin.set("typeface", SANS)
    theme_part._blob = etree.tostring(root, xml_declaration=True,
                                      encoding="UTF-8", standalone=True)


# ---- small XML helpers -------------------------------------------------------
def _remove_ph(shapes_owner, types=("dt", "ftr", "sldNum")):
    """Drop date/footer/slide-number placeholders from a master or layout."""
    spTree = shapes_owner.shapes._spTree
    for sp in list(spTree.findall(qn("p:sp"))):
        ph = sp.find(f".//{qn('p:ph')}")
        if ph is not None and ph.get("type") in types:
            spTree.remove(sp)


def _find_ph(owner, ph_type=None, idx=None):
    for ph in owner.placeholders:
        el = ph._element.find(f".//{qn('p:ph')}")
        if ph_type is not None and el.get("type") != ph_type:
            continue
        if idx is not None and el.get("idx", "0") != str(idx):
            continue
        return ph
    return None


def _set_lst_style(ph, lvl_xml):
    """Replace the placeholder's <a:lstStyle> with the given lvlNpPr markup."""
    txBody = ph._element.find(qn("p:txBody"))
    old = txBody.find(qn("a:lstStyle"))
    if old is not None:
        txBody.remove(old)
    lst = parse_xml(f"<a:lstStyle {A}>{lvl_xml}</a:lstStyle>")
    txBody.insert(list(txBody).index(txBody.find(qn("a:bodyPr"))) + 1, lst)


def _set_prompt(ph, text):
    """Set the layout placeholder's prompt text (shown as 'Click to add ...')."""
    txBody = ph._element.find(qn("p:txBody"))
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    p = parse_xml(
        f"<a:p {A}><a:r><a:rPr lang='en-US'/><a:t>{text}</a:t></a:r></a:p>"
    )
    txBody.append(p)


def _place(ph, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    ph.left, ph.top, ph.width, ph.height = int(left), int(top), int(width), int(height)
    ph.text_frame.word_wrap = True
    ph.text_frame.vertical_anchor = anchor


def _next_id(spTree):
    ids = [int(e.get("id")) for e in spTree.iter(qn("p:cNvPr")) if e.get("id")]
    return max(ids, default=1) + 1


def _lvl(n, *, sz, color, bold=False, font=None, caps=False, align=None,
         mar_l=0, indent=0, space_after=0, bullet=None, bullet_color=None):
    """Build one <a:lvlNpPr> override string."""
    attrs = f'marL="{mar_l}" indent="{indent}"'
    if align:
        attrs += f' algn="{align}"'
    inner = ""
    if space_after:
        inner += f'<a:spcAft><a:spcPts val="{space_after}"/></a:spcAft>'
    if bullet is None:
        inner += "<a:buNone/>"
    else:
        inner += (f'<a:buClr><a:schemeClr val="{bullet_color}"/></a:buClr>'
                  f'<a:buFont typeface="Arial"/><a:buChar char="{bullet}"/>')
    rpr = f'sz="{sz}"'
    if bold:
        rpr += ' b="1"'
    if caps:
        rpr += ' cap="all"'
    latin = f'<a:latin typeface="{font}"/>' if font else '<a:latin typeface="+mn-lt"/>'
    inner += (f'<a:defRPr {rpr}><a:solidFill><a:schemeClr val="{color}"/>'
              f"</a:solidFill>{latin}</a:defRPr>")
    return f"<a:lvl{n}pPr {attrs}>{inner}</a:lvl{n}pPr>"


def add_eyebrow(layout):
    """Add the small mono eyebrow placeholder (body idx=13) to a layout."""
    spTree = layout.shapes._spTree
    sp = parse_xml(f"""
<p:sp {nsdecls('p')} {A}>
  <p:nvSpPr>
    <p:cNvPr id="{_next_id(spTree)}" name="Eyebrow Placeholder"/>
    <p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
    <p:nvPr><p:ph type="body" sz="quarter" idx="{EYEBROW_IDX}"/></p:nvPr>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{int(M)}" y="{int(Inches(0.35))}"/>
            <a:ext cx="{int(USABLE)}" cy="{int(Inches(0.4))}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
  <p:txBody>
    <a:bodyPr wrap="square" anchor="t"/>
    <a:lstStyle>{_lvl(1, sz=1300, color="accent4", bold=True, font=MONO, caps=True)}</a:lstStyle>
    <a:p><a:r><a:rPr lang="en-US"/><a:t>MSACL · DS301 Deep Learning · Segment n · Lecture n</a:t></a:r></a:p>
  </p:txBody>
</p:sp>""")
    spTree.append(sp)


def add_rect(layout, name, left, top, width, height, scheme_color):
    """Add a static (non-placeholder) filled rectangle to a layout."""
    spTree = layout.shapes._spTree
    sp = parse_xml(f"""
<p:sp {nsdecls('p')} {A}>
  <p:nvSpPr>
    <p:cNvPr id="{_next_id(spTree)}" name="{name}"/>
    <p:cNvSpPr/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{int(left)}" y="{int(top)}"/>
            <a:ext cx="{int(width)}" cy="{int(height)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill><a:schemeClr val="{scheme_color}"/></a:solidFill>
    <a:ln><a:noFill/></a:ln>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>""")
    spTree.append(sp)


# ---- master ------------------------------------------------------------------
MASTER_TXSTYLES = f"""
<p:txStyles {nsdecls('p')} {A}>
  <p:titleStyle>
    {_lvl(1, sz=3400, color="tx1", bold=True)}
  </p:titleStyle>
  <p:bodyStyle>
    {_lvl(1, sz=2600, color="tx1", mar_l=274638, indent=-274638,
          space_after=1800, bullet="•", bullet_color="accent1")}
    {_lvl(2, sz=2100, color="accent4", mar_l=549275, indent=-274638,
          space_after=1400)}
    {_lvl(3, sz=1800, color="tx2", mar_l=823913, indent=-274638,
          space_after=1200, bullet="–", bullet_color="accent4")}
  </p:bodyStyle>
  <p:otherStyle>
    {_lvl(1, sz=1800, color="tx1")}
  </p:otherStyle>
</p:txStyles>
"""


def restyle_master(master):
    master.element.find(qn("p:cSld")).set("name", "MSACL DS301")
    bg = master.background
    bg.fill.solid()
    # bg1 = theme lt1 = PAPER; recoloring the theme recolors every deck ground
    from pptx.enum.dml import MSO_THEME_COLOR
    bg.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
    _remove_ph(master)
    old = master.element.find(qn("p:txStyles"))
    master.element.replace(old, parse_xml(MASTER_TXSTYLES))
    title = _find_ph(master, "title")
    _place(title, M, Inches(0.8), USABLE, Inches(1.1))
    body = _find_ph(master, "body")
    _place(body, M, Inches(2.1), USABLE, Inches(4.85))


# ---- layouts -----------------------------------------------------------------
def config_title_slide(layout):
    title = _find_ph(layout, "ctrTitle")
    _place(title, M, Inches(2.5), USABLE, Inches(1.5))
    _set_lst_style(title, _lvl(1, sz=4400, color="tx1", bold=True, align="l"))
    sub = _find_ph(layout, "subTitle")
    _place(sub, M, Inches(4.2), USABLE, Inches(1.0))
    _set_lst_style(sub, _lvl(1, sz=2200, color="accent1", align="l"))
    _set_prompt(sub, "Click to add a one-line subtitle")
    add_rect(layout, "Amber bar", M, Inches(2.25), Inches(1.4), Inches(0.06), "accent2")


def config_title_and_body(layout):
    _place(_find_ph(layout, "title"), M, Inches(0.8), USABLE, Inches(1.1))
    body = _find_ph(layout, idx=1)
    _place(body, M, Inches(2.1), USABLE, Inches(4.85))


def config_two_column(layout):
    _place(_find_ph(layout, "title"), M, Inches(0.8), USABLE, Inches(1.1))
    col_w = Inches(5.87)
    left = _find_ph(layout, idx=1)
    _place(left, M, Inches(2.1), col_w, Inches(4.85))
    _set_lst_style(left, _lvl(1, sz=2200, color="tx1", mar_l=274638,
                              indent=-274638, space_after=1400,
                              bullet="•", bullet_color="accent1"))
    right = _find_ph(layout, idx=2)
    _place(right, Inches(6.87), Inches(2.1), col_w, Inches(4.85))
    _set_lst_style(right, _lvl(1, sz=2200, color="tx1", mar_l=274638,
                               indent=-274638, space_after=1400,
                               bullet="•", bullet_color="accent1"))


def config_screenshot(layout):
    _place(_find_ph(layout, "title"), M, Inches(0.8), USABLE, Inches(1.1))
    pic = _find_ph(layout, "pic")
    _place(pic, M, Inches(2.0), USABLE, Inches(4.35), MSO_ANCHOR.MIDDLE)
    cap = _find_ph(layout, idx=2)
    _place(cap, M, Inches(6.5), USABLE, Inches(0.5))
    _set_lst_style(cap, _lvl(1, sz=1300, color="accent4", align="ctr"))
    _set_prompt(cap, "Source · license — click to add caption")


def config_title_only(layout):
    _place(_find_ph(layout, "title"), M, Inches(0.8), USABLE, Inches(1.1))


LAYOUT_CONFIG = {
    "Title Slide": config_title_slide,
    "Title and Body": config_title_and_body,
    "Two Column": config_two_column,
    "Screenshot + Caption": config_screenshot,
    "Title Only": config_title_only,
    "Blank": None,
}


def prune_and_configure_layouts(master):
    # drop the layouts we don't keep: remove from sldLayoutIdLst + master rels
    id_lst = master.element.find(qn("p:sldLayoutIdLst"))
    rels = master.part.rels
    for entry in list(id_lst):
        rId = entry.get(qn("r:id"))
        layout_part = rels[rId].target_part
        name = layout_part.slide_layout.name
        if name not in KEEP:
            id_lst.remove(entry)
            rels.pop(rId)
    # rename + configure the keepers
    for layout in master.slide_layouts:
        new_name = KEEP[layout.name]
        layout.element.find(qn("p:cSld")).set("name", new_name)
        _remove_ph(layout)
        if new_name != "Blank":
            add_eyebrow(layout)
        cfg = LAYOUT_CONFIG[new_name]
        if cfg:
            cfg(layout)


def normalize_master_rels(master):
    """Renumber master rels contiguously (rId1..rIdN layouts, then theme) and
    rewrite sldLayoutIdLst r:id to match, so adopt_template.py can replicate
    the rels on a fresh part and get identical rIds."""
    rels = master.part.rels
    id_lst = master.element.find(qn("p:sldLayoutIdLst"))
    # put the gallery in teaching order before renumbering
    want = list(KEEP.values())
    entries = sorted(
        id_lst,
        key=lambda e: want.index(rels[e.get(qn("r:id"))].target_part.slide_layout.name),
    )
    for e in entries:
        id_lst.append(e)  # re-append moves the element to the end, in order
    ordered = [(e, rels[e.get(qn("r:id"))].target_part) for e in id_lst]
    theme = master.part.part_related_by(RT.THEME)
    old_keys = list(rels.keys())
    for k in old_keys:
        rels.pop(k)
    for entry, layout_part in ordered:
        new_rId = rels.get_or_add(RT.SLIDE_LAYOUT, layout_part)
        entry.set(qn("r:id"), new_rId)
    rels.get_or_add(RT.THEME, theme)


# ------------------------------------------------------------------------------
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = EMU_W, EMU_H
    master = prs.slide_masters[0]
    restyle_theme(master.part.part_related_by(RT.THEME))
    restyle_master(master)
    prune_and_configure_layouts(master)
    normalize_master_rels(master)
    cp = prs.core_properties
    cp.title = "MSACL DS301 Deep Learning — slide template"
    cp.category = "MSACL DS301 Deep Learning"
    cp.comments = (
        f"Course template; edit master/layouts here, then propagate with "
        f"tools/adopt_template.py. Fonts {SANS}/{MONO}. "
        f"paper#{PAPER} ink#{INK} teal#{SIGNAL} amber#{ROI} red#{NEG}"
    )
    TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(TEMPLATE))
    return TEMPLATE


def verify():
    prs = Presentation(str(TEMPLATE))
    master = prs.slide_masters[0]
    names = [l.name for l in master.slide_layouts]
    assert names == list(KEEP.values()), names
    assert len(prs.slides) == 0
    for layout in master.slide_layouts:
        phs = [(p.placeholder_format.type, p.placeholder_format.idx)
               for p in layout.placeholders]
        print(f"  {layout.name}: {phs}")
    print(f"OK: {TEMPLATE.relative_to(ROOT)} "
          f"({len(names)} layouts, {int(prs.slide_width)/914400:.3f}x"
          f"{int(prs.slide_height)/914400}in)")


def main():
    build()
    verify()
    return 0


if __name__ == "__main__":
    sys.exit(main())
