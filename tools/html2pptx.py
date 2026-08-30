#!/usr/bin/env python3
"""Convert HTML slide decks to PowerPoint.

Deck contract (see CLAUDE.md):
  - one element matching ".slide" per slide, designed at 1280x720
  - optional speaker notes in a child <aside class="notes"> (hidden during
    rendering, copied into the PowerPoint notes pane)
  - decks must be self-contained (inline CSS or slides/assets/) so headless
    rendering works offline

Each slide is screenshotted at 2x scale with headless Chromium and placed
full-bleed on a 16:9 PowerPoint slide. Trade-off: PPTX slides are images —
pixel-perfect but not text-editable; the HTML remains the editable source.

Usage:
    python tools/html2pptx.py slides/html/lecture01_intro.html
    python tools/html2pptx.py --all

Setup: pip install -r tools/requirements.txt && playwright install chromium
"""

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HTML_DIR = ROOT / "slides" / "html"
PPTX_DIR = ROOT / "slides" / "pptx"

HIDE_NOTES_CSS = "aside.notes { display: none !important; }"


def convert(deck_path: Path, page, tmpdir: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    page.goto(deck_path.resolve().as_uri())
    page.add_style_tag(content=HIDE_NOTES_CSS)
    page.wait_for_load_state("networkidle")

    slides = page.locator(".slide")
    count = slides.count()
    if count == 0:
        raise ValueError(f"{deck_path.name}: no elements match '.slide'")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for i in range(count):
        png = tmpdir / f"slide_{i:03d}.png"
        slides.nth(i).screenshot(path=str(png))
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(png), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        notes = slides.nth(i).locator("aside.notes")
        if notes.count():
            text = notes.first.evaluate("el => el.innerText").strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text

    PPTX_DIR.mkdir(parents=True, exist_ok=True)
    dest = PPTX_DIR / (deck_path.stem + ".pptx")
    prs.save(str(dest))
    print(f"{deck_path.name}: {count} slide(s) -> {dest.relative_to(ROOT)}")
    return dest


def main(argv):
    if len(argv) < 2 or argv[1] in ("-h", "--help"):
        print(__doc__)
        return 0
    if argv[1] == "--all":
        targets = sorted(HTML_DIR.glob("*.html"))
        if not targets:
            print(f"No decks found in {HTML_DIR}")
            return 1
    else:
        targets = [Path(a).resolve() for a in argv[1:]]

    from playwright.sync_api import sync_playwright

    status = 0
    with sync_playwright() as pw, tempfile.TemporaryDirectory() as tmp:
        browser = pw.chromium.launch()
        page = browser.new_page(
            viewport={"width": 1280, "height": 720}, device_scale_factor=2
        )
        for deck in targets:
            try:
                convert(deck, page, Path(tmp))
            except Exception as exc:  # keep converting the rest of the decks
                print(f"ERROR in {deck.name}: {exc}", file=sys.stderr)
                status = 1
        browser.close()
    return status


if __name__ == "__main__":
    sys.exit(main(sys.argv))
