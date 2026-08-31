#!/usr/bin/env python3
"""Sanity-check generated PPTX decks for TEXT that overflows its box.

python-pptx cannot measure rendered text and there is no LibreOffice on the
build host, so tools/spec2pptx.py estimates wrapped text height deterministically
and sizes boxes to fit. This checker independently re-estimates, for every
text-bearing shape, how tall the wrapped text needs to be and FLAGS any shape
whose estimated text height exceeds its set box height by more than TOL. It uses
the SAME estimator constants as spec2pptx.py (ADVANCE, LH). Non-wrapping
single-line pills (amber annotation labels) are skipped because they cannot
overflow via wrapping.

Usage:
    python tools/check_text_overflow.py slides/pptx/lecture04_training_practice.pptx
    python tools/check_text_overflow.py --all      # every deck in slides/pptx/
Exit code is nonzero if any overflow is found, so it can gate a build.
"""
import sys
from pathlib import Path

from pptx import Presentation

EMU = 914400.0
ADVANCE = 0.55      # avg glyph advance as fraction of point size (Avenir Next-ish)
LH = 1.2            # line-height multiple
TOL = 0.06          # in slack before we call it overflow

ROOT = Path(__file__).resolve().parent.parent
PPTX_DIR = ROOT / "slides" / "pptx"


def est_lines(text, usable_in, font_pt):
    if not text:
        return 0
    cw = ADVANCE * font_pt / 72.0
    cpl = max(usable_in / cw, 1.0)
    lines = 0
    for para in text.split("\n"):
        words = para.split(" ")
        cur = 0
        pl = 1
        for w in words:
            add = (len(w) + 1) if cur > 0 else len(w)
            if cur + add > cpl and cur > 0:
                pl += 1
                cur = len(w)
            else:
                cur += add
        lines += pl
    return lines


def est_text_height(shape):
    tf = shape.text_frame
    w = (shape.width or 0) / EMU
    ml = (tf.margin_left / EMU) if tf.margin_left is not None else 0.1
    mr = (tf.margin_right / EMU) if tf.margin_right is not None else 0.1
    mt = (tf.margin_top / EMU) if tf.margin_top is not None else 0.05
    mb = (tf.margin_bottom / EMU) if tf.margin_bottom is not None else 0.05
    usable = max(w - ml - mr, 0.5)
    total_h = mt + mb
    for para in tf.paragraphs:
        runs = para.runs
        if not runs:
            continue
        ptext = "".join(r.text for r in runs)
        fsizes = [r.font.size.pt for r in runs if r.font.size]
        fpt = max(fsizes) if fsizes else 18.0
        lines = est_lines(ptext, usable, fpt)
        lh_in = fpt * LH / 72.0
        total_h += lines * lh_in
        sa = (para.space_after.pt / 72.0) if para.space_after else 0
        sb = (para.space_before.pt / 72.0) if para.space_before else 0
        total_h += sa + sb
    return total_h


def check_deck(path):
    prs = Presentation(str(path))
    over = []
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False):
                continue
            if not (sh.text_frame.text or "").strip():
                continue
            # Non-wrapping single-line pills (amber annotation labels) can't
            # overflow via wrapping — skip them.
            if sh.text_frame.word_wrap is False:
                continue
            box_h = (sh.height or 0) / EMU
            est = est_text_height(sh)
            if est > box_h + TOL:
                over.append((i + 1, box_h, est, sh.text_frame.text[:60]))
    for (sl, bh, est, txt) in over:
        print(f"  OVERFLOW {path.name} slide {sl}: box={bh:.2f}in est={est:.2f}in :: {txt!r}")
    print(f"[{'OVERFLOW' if over else 'ok'}] {path.name}: {len(over)} overflowing box(es)")
    return len(over)


def main(argv):
    args = argv[1:]
    if not args or args[0] == "--all":
        paths = sorted(p for p in PPTX_DIR.glob("*.pptx") if not p.name.startswith("~$"))
    else:
        paths = [Path(a) for a in args]
    total = 0
    for p in paths:
        total += check_deck(p)
    print(f"\n{'OVERFLOW' if total else 'PASS'}: {total} overflowing text box(es) across {len(paths)} deck(s)")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
