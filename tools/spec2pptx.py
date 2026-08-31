#!/usr/bin/env python3
"""Build NATIVE, editable PowerPoint decks from a compact YAML spec.

Why this exists (replaces the html->screenshot pipeline):
  - the old tools/html2pptx.py rendered each HTML slide to a PNG and pasted it
    full-bleed, so the .pptx was flat images: not editable, web-styled, and
    every deck was 700-900 lines of hand-written HTML/CSS/SVG.
  - this builder reads a small YAML deck spec (~10-15 lines per slide) and emits
    real PowerPoint objects: native text boxes, tables, shapes, connectors and
    placed images. Titles/body stay editable in PowerPoint; look is driven by a
    shared theme here, not CSS. Real licensed images live under slides/assets/img
    and are placed with `image:`; hand-drawn SVG is used only where rule 6 (the
    numeric mechanics) truly needs it.

Spec contract (see slides/spec/*.yaml for a worked example):

    deck:
      title: "Lecture 5 ..."           # core doc title
      eyebrow: "MSACL · DS301 ..."     # default eyebrow shown top-left
    slides:
      - type: title
        eyebrow: "..."                  # overrides deck.eyebrow
        title: "Big title"
        subtitle: "one line"
        notes: "speaker notes"

      - type: bullets
        eyebrow: "..."
        title: "Slide title"
        bullets:                        # str, or {text, time, sub}
          - text: "Why CNNs fit spectra"
            time: "8 min"
        callout: {kind: ms, text: "..."}   # kind: ms|check|pair|plain
        notes: "..."

      - type: two-col
        title: "..."
        left:  {image: img/sparrow.jpg, caption: "photo: ..., CC BY-SA 2.0"}
        right: {lead: ["para 1", "para 2"], callout: {kind: ms, text: "..."}}
        notes: "..."

      - type: image                    # one big image + caption + optional callout
        title: "..."
        image: img/vgg16.png
        img_h: 4.2                       # optional: taller image box (default 3.5in)
        caption: "VGG-16 · CC BY-SA 4.0"
        calc: "every filter is 3x3 ..."   # mono strip under the figure
        callout: {kind: plain, text: "..."}

      - type: pipeline                 # stage boxes joined by arrows
        title: "The whole CNN"
        stages:                        # {label, sub, kind: plain|conv|pool|dashed}
          - {label: input, sub: spectrum}
          - {label: conv, sub: pattern maps, kind: conv}
        callout: {kind: plain, text: "..."}

      - type: grid                     # numeric table (conv/pooling mechanics, rule 6)
        title: "..."
        table:
          data: [[1,0,0],[0,1,0]]
          highlight: [[0,0],[1,1]]     # cells drawn in amber (the sliding window)
        calc: "3x3 filter, stride 1 -> 4x4 map"

Usage:
    python tools/spec2pptx.py slides/spec/lecture05_cnns.yaml
    python tools/spec2pptx.py --all
"""

import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SPEC_DIR = ROOT / "slides" / "spec"
PPTX_DIR = ROOT / "slides" / "pptx"
ASSETS = ROOT / "slides" / "assets"

# ---- theme (mirrors slides/assets/theme.css, but native) --------------------
PAPER = RGBColor(0xFB, 0xFB, 0xF8)
INK = RGBColor(0x1B, 0x20, 0x25)
INK_SOFT = RGBColor(0x3D, 0x44, 0x4C)
MUTED = RGBColor(0x8A, 0x90, 0x99)
HAIRLINE = RGBColor(0xD9, 0xD9, 0xD2)
SIGNAL = RGBColor(0x0E, 0x7C, 0x7B)
SIGNAL_SOFT = RGBColor(0xE3, 0xF0, 0xEF)
ROI = RGBColor(0xE0, 0x9E, 0x2F)
ROI_SOFT = RGBColor(0xFB, 0xF0, 0xDC)
ROI_INK = RGBColor(0xA9, 0x72, 0x1A)
NEG = RGBColor(0xC4, 0x53, 0x4F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "Avenir Next"
MONO = "Menlo"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)

CALLOUT_STYLE = {  # kind -> (bar color, prefix)
    "ms": (SIGNAL, "MASS SPEC · "),
    "check": (ROI, "CHECKPOINT · "),
    "pair": (ROI, "PAIR · "),
    "plain": (INK_SOFT, ""),
}


def _set(run, size, color, bold=False, italic=False, font=SANS):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def _bg(slide, prs):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = PAPER
    r.line.fill.background()
    r.shadow.inherit = False
    slide.shapes._spTree.remove(r._element)
    slide.shapes._spTree.insert(2, r._element)
    return r


def _text(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def _eyebrow(slide, text):
    if not text:
        return
    _, tf = _text(slide, MARGIN, Inches(0.35), Inches(12), Inches(0.4))
    r = tf.paragraphs[0].add_run()
    r.text = text.upper()
    _set(r, 13, MUTED, bold=True)
    tf.paragraphs[0].runs[0].font.name = MONO


def _title(slide, text, top=Inches(0.8)):
    _, tf = _text(slide, MARGIN, top, Inches(12.1), Inches(1.1))
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    _set(r, 34, INK, bold=True)


def _callout(slide, spec, top, width=Inches(11.5), left=None):
    if not spec:
        return
    kind = spec.get("kind", "plain")
    bar, prefix = CALLOUT_STYLE.get(kind, CALLOUT_STYLE["plain"])
    left = left if left is not None else MARGIN
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(1.0))
    box.fill.solid()
    box.fill.fore_color.rgb = SIGNAL_SOFT if kind == "ms" else RGBColor(0xF1, 0xF1, 0xEC)
    box.line.color.rgb = bar
    box.line.width = Pt(1)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    if prefix:
        rp = p.add_run()
        rp.text = prefix
        _set(rp, 12, bar, bold=True, font=MONO)
    r = p.add_run()
    r.text = spec["text"]
    _set(r, 17, INK_SOFT)
    return box


# ---- slide builders ---------------------------------------------------------
def slide_title(slide, prs, s):
    _eyebrow(slide, s.get("eyebrow"))
    _, tf = _text(slide, MARGIN, Inches(2.4), Inches(12.1), Inches(2.2), MSO_ANCHOR.MIDDLE)
    r = tf.paragraphs[0].add_run()
    r.text = s["title"]
    _set(r, 44, INK, bold=True)
    if s.get("subtitle"):
        p = tf.add_paragraph()
        p.space_before = Pt(16)
        r = p.add_run()
        r.text = s["subtitle"]
        _set(r, 22, SIGNAL)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(2.25), Inches(1.4), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ROI
    bar.line.fill.background()
    bar.shadow.inherit = False


def _bullets(tf, bullets, size=26):
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(18)
        if isinstance(b, str):
            b = {"text": b}
        r = p.add_run()
        r.text = "•  "
        _set(r, size, SIGNAL, bold=True)
        r = p.add_run()
        r.text = b["text"]
        _set(r, size, INK)
        if b.get("time"):
            r = p.add_run()
            r.text = f"   · {b['time']}"
            _set(r, size - 5, MUTED, font=MONO)
        if b.get("sub"):
            sp = tf.add_paragraph()
            sp.space_after = Pt(18)
            r = sp.add_run()
            r.text = "     " + b["sub"]
            _set(r, size - 5, MUTED)


def slide_bullets(slide, prs, s):
    _eyebrow(slide, s.get("eyebrow"))
    _title(slide, s["title"])
    _, tf = _text(slide, MARGIN, Inches(2.1), Inches(11.8), Inches(3.6))
    _bullets(tf, s.get("bullets", []))
    if s.get("callout"):
        _callout(slide, s["callout"], Inches(6.1))


def _annotate(slide, img_x, img_y, img_w, img_h, boxes):
    """Draw amber bounding boxes (and optional labels) over a placed image.
    Each box: {x, y, w, h} as fractions [0-1] of the image, optional 'label'
    plus optional label anchor 'lx','ly' (fractions); label defaults above box."""
    for b in boxes:
        bx = int(img_x + b["x"] * img_w)
        by = int(img_y + b["y"] * img_h)
        bw = int(b["w"] * img_w)
        bh = int(b["h"] * img_h)
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, bh)
        rect.fill.background()
        rect.line.color.rgb = ROI
        rect.line.width = Pt(2.5)
        rect.shadow.inherit = False
        if b.get("label"):
            lx = int(img_x + b.get("lx", b["x"]) * img_w)
            ly = int(img_y + b.get("ly", max(b["y"] - 0.09, 0)) * img_h)
            lab = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, lx, ly, Inches(1.9), Inches(0.34)
            )
            lab.fill.solid()
            lab.fill.fore_color.rgb = WHITE
            lab.line.color.rgb = ROI
            lab.line.width = Pt(0.75)
            lab.shadow.inherit = False
            tf = lab.text_frame
            tf.margin_top = tf.margin_bottom = Inches(0.02)
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = b["label"]
            _set(r, 12, ROI_INK, bold=True, font=MONO)


def _place_image(slide, path, left, top, max_w, max_h, caption=None, boxes=None):
    """Place an image scaled into (max_w, max_h) and return the Y of its bottom
    edge (below the caption, if any) so callers can flow more content beneath."""
    img = ASSETS / path
    from PIL import Image

    with Image.open(img) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    x = int(left + (max_w - w) // 2)
    slide.shapes.add_picture(str(img), x, int(top), width=w, height=h)
    if boxes:
        _annotate(slide, x, int(top), w, h, boxes)
    bottom = int(top) + h
    if caption:
        cap_top = bottom + Inches(0.06)
        cap_h = Inches(0.45)
        _, tf = _text(slide, left, cap_top, max_w, cap_h)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = caption
        _set(r, 12, MUTED, italic=True)
        bottom = cap_top + cap_h
    return bottom


def slide_two_col(slide, prs, s):
    _eyebrow(slide, s.get("eyebrow"))
    _title(slide, s["title"])
    colw = Inches(5.9)
    left = s.get("left", {})
    right = s.get("right", {})
    if left.get("image"):
        _place_image(slide, left["image"], MARGIN, Inches(2.15), int(colw), Inches(3.2), left.get("caption"), left.get("boxes"))
    elif left.get("lead"):
        _, tf = _text(slide, MARGIN, Inches(2.2), colw, Inches(4))
        for i, para in enumerate(left["lead"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            r = p.add_run()
            r.text = para
            _set(r, 22, INK)
    rx = Inches(6.9)
    if right.get("image"):
        _place_image(slide, right["image"], rx, Inches(2.15), int(colw), Inches(3.2), right.get("caption"), right.get("boxes"))
    else:
        _, tf = _text(slide, rx, Inches(2.2), colw, Inches(3.4))
        for i, para in enumerate(right.get("lead", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            r = p.add_run()
            r.text = para
            _set(r, 22, INK)
    if right.get("callout"):
        _callout(slide, right["callout"], Inches(5.9), width=colw, left=rx)
    elif s.get("callout"):
        _callout(slide, s["callout"], Inches(6.2))


def slide_image(slide, prs, s):
    _eyebrow(slide, s.get("eyebrow"))
    _title(slide, s["title"])
    # img_h (optional) lets a tall/near-square figure use more vertical space;
    # default 3.5in keeps every existing deck (e.g. lecture05) byte-compatible.
    max_h = Inches(s.get("img_h", 3.5))
    y = _place_image(slide, s["image"], MARGIN, Inches(2.0), Inches(12.1), max_h, s.get("caption"), s.get("boxes"))
    y += Inches(0.1)
    if s.get("calc"):
        _, tf = _text(slide, MARGIN, y, Inches(12.1), Inches(0.5))
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = s["calc"]
        _set(r, 15, INK_SOFT, font=MONO)
        y += Inches(0.55)
    if s.get("callout"):
        _callout(slide, s["callout"], min(y, Inches(6.5)))


def slide_pipeline(slide, prs, s):
    _eyebrow(slide, s.get("eyebrow"))
    _title(slide, s["title"])
    stages = s["stages"]
    n = len(stages)
    gap = Inches(0.25)
    total = Inches(12.1)
    box_w = int((total - gap * (n - 1)) / n)
    y = Inches(3.0)
    box_h = Inches(1.5)
    kind_color = {"conv": SIGNAL, "pool": ROI, "dashed": NEG, "plain": INK_SOFT}
    for i, st in enumerate(stages):
        x = int(MARGIN + i * (box_w + gap))
        kind = st.get("kind", "plain")
        col = kind_color.get(kind, INK_SOFT)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = col
        box.line.width = Pt(2.5 if kind != "plain" else 1.5)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = st["label"]
        _set(r, 18, col if kind != "plain" else INK, bold=True)
        if st.get("sub"):
            sp = tf.add_paragraph()
            sp.alignment = PP_ALIGN.CENTER
            r = sp.add_run()
            r.text = st["sub"]
            _set(r, 12, MUTED)
        if i < n - 1:
            ax = int(x + box_w)
            cy = int(y + box_h / 2)
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, cy, int(ax + gap), cy)
            conn.line.color.rgb = MUTED
            conn.line.width = Pt(1.5)
    if s.get("callout"):
        _callout(slide, s["callout"], Inches(5.4))


def slide_grid(slide, prs, s):
    _eyebrow(slide, s.get("eyebrow"))
    _title(slide, s["title"])
    spec = s["table"]
    data = spec["data"]
    highlight = {tuple(c) for c in spec.get("highlight", [])}
    rows, cols = len(data), len(data[0])
    cell = Inches(0.62)
    tw, th = cell * cols, cell * rows
    x0 = (EMU_W - tw) // 2
    y0 = Inches(2.4)
    gt = slide.shapes.add_table(rows, cols, x0, y0, tw, th).table
    for c in range(cols):
        gt.columns[c].width = cell
    for r_ in range(rows):
        gt.rows[r_].height = cell
    for r_ in range(rows):
        for c in range(cols):
            cellobj = gt.cell(r_, c)
            cellobj.text = str(data[r_][c])
            para = cellobj.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            _set(run, 20, ROI_INK if (r_, c) in highlight else INK, bold=(r_, c) in highlight, font=MONO)
            cellobj.fill.solid()
            cellobj.fill.fore_color.rgb = ROI_SOFT if (r_, c) in highlight else WHITE
            cellobj.vertical_anchor = MSO_ANCHOR.MIDDLE
    if s.get("calc"):
        _, tf = _text(slide, MARGIN, y0 + th + Inches(0.3), Inches(12.1), Inches(0.6))
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = s["calc"]
        _set(r, 16, INK_SOFT, font=MONO)
    if s.get("callout"):
        _callout(slide, s["callout"], Inches(6.2))


BUILDERS = {
    "title": slide_title,
    "bullets": slide_bullets,
    "two-col": slide_two_col,
    "image": slide_image,
    "pipeline": slide_pipeline,
    "grid": slide_grid,
}


def build(spec_path: Path) -> Path:
    spec = yaml.safe_load(spec_path.read_text())
    deck = spec.get("deck", {})
    default_eyebrow = deck.get("eyebrow", "")

    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    blank = prs.slide_layouts[6]

    for s in spec["slides"]:
        slide = prs.slides.add_slide(blank)
        _bg(slide, prs)
        s.setdefault("eyebrow", default_eyebrow)
        builder = BUILDERS.get(s["type"])
        if not builder:
            raise ValueError(f"unknown slide type: {s['type']}")
        builder(slide, prs, s)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    if deck.get("title"):
        prs.core_properties.title = deck["title"]
    PPTX_DIR.mkdir(parents=True, exist_ok=True)
    dest = PPTX_DIR / (spec_path.stem + ".pptx")
    prs.save(str(dest))
    print(f"{spec_path.name}: {len(spec['slides'])} slide(s) -> {dest.relative_to(ROOT)}")
    return dest


def main(argv):
    if len(argv) < 2 or argv[1] in ("-h", "--help"):
        print(__doc__)
        return 0
    targets = sorted(SPEC_DIR.glob("*.yaml")) if argv[1] == "--all" else [Path(a).resolve() for a in argv[1:]]
    status = 0
    for spec in targets:
        try:
            build(spec)
        except Exception as exc:
            print(f"ERROR in {spec.name}: {exc}", file=sys.stderr)
            status = 1
    return status


if __name__ == "__main__":
    sys.exit(main(sys.argv))
